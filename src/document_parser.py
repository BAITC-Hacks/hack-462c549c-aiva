"""Format dispatch layer. All supported inputs become ParsedDocument."""

import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from .models import Fragment, FunctionRecord, ParsedDocument
from .parser import extract_functions, extract_units, parse_pdf

CLAUSE_RE = re.compile(r"(?<![\w.])(\d+(?:\.\d+){1,3})(?:\.)?\s+")


def parse_document(path: str | Path) -> ParsedDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".xlsx":
        return parse_xlsx(path)
    if suffix == ".xls":
        raise ValueError("Формат .xls не поддерживается. Используйте .xlsx, .docx или .pdf.")
    raise ValueError(f"Неподдерживаемый формат: {path.suffix or 'без расширения'}")


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _fragment(document_name, index, raw, locator, clause_id=None):
    cleaned = _clean(raw)
    return Fragment(
        fragment_id=f"{Path(document_name).stem}-f{index}",
        document_name=Path(document_name).name,
        page_number=None,
        clause_id=clause_id,
        raw_text=raw,
        cleaned_text=cleaned,
        source_locator=locator,
    )


def parse_docx(path: str | Path) -> ParsedDocument:
    path = Path(path)
    document = Document(str(path))
    fragments = []
    index = 0
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        raw = paragraph.text or ""
        if not _clean(raw):
            continue
        index += 1
        match = CLAUSE_RE.search(_clean(raw))
        fragments.append(_fragment(path.name, index, raw, f"paragraph {paragraph_number}", match.group(1) if match else None))
    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            raw = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if not _clean(raw):
                continue
            index += 1
            match = CLAUSE_RE.search(_clean(raw))
            fragments.append(_fragment(path.name, index, raw, f"table {table_number}, row {row_number}", match.group(1) if match else None))
    return _build_parsed(path, fragments)


def parse_xlsx(path: str | Path) -> ParsedDocument:
    path = Path(path)
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    fragments = []
    index = 0
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(), start=1):
            values = [str(cell.value).strip() for cell in row if cell.value is not None and str(cell.value).strip()]
            if not values:
                continue
            raw = " | ".join(values)
            index += 1
            first = row[0].coordinate
            last = row[-1].coordinate
            match = CLAUSE_RE.search(_clean(raw))
            fragments.append(_fragment(path.name, index, raw, f"sheet {sheet.title}, row {row_number}, cells {first}:{last}", match.group(1) if match else None))
    return _build_parsed(path, fragments)


def _build_parsed(path, fragments):
    clauses = {}
    for fragment in fragments:
        if fragment.clause_id:
            clauses.setdefault(fragment.clause_id, fragment)
    return ParsedDocument(
        document_name=path.name,
        fragments=fragments,
        clauses=clauses,
        units=extract_units(fragments),
        functions=extract_functions(fragments),
    )
