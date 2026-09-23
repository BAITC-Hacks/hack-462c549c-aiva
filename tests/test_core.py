from pathlib import Path

from src.analyzer import compare_documents
from src.models import Fragment, ParsedDocument
from src.parser import parse_pdf
from src.document_parser import parse_document, parse_docx, parse_xlsx
from src.semantic import enhance_rows, validate_result
import src.semantic as semantic
from src.report import build_recommendations, expert_conclusion, format_clause_ids, group_recommendations, management_conclusion, review_key, review_priority
from src.decision_workflow import apply_employee_decision, decision_record, revalidate_records, russian_impact, russian_recommendation
from src.ui_state import get_analysis_state, save_analysis_state, save_human_decision
from app import diff_html, source_label, restore_uploaded_name


SAMPLE = Path("data/sample")


def test_sample_pdf_extracts_text_and_clauses():
    path = next(SAMPLE.glob("*.pdf"))
    parsed = parse_pdf(path)
    assert parsed.fragments
    assert parsed.clauses
    assert any(fragment.clause_id == "3.4" for fragment in parsed.fragments)


def test_source_fragments_exist_for_sample_units():
    path = next(SAMPLE.glob("*.pdf"))
    parsed = parse_pdf(path)
    ids = {fragment.fragment_id for fragment in parsed.fragments}
    assert all(unit["fragment_id"] in ids for unit in parsed.units)


def make_doc(name, clauses):
    fragments = [Fragment(fragment_id=f"p1-c{k}", document_name=name, page_number=1, clause_id=k, raw_text=v, cleaned_text=v) for k, v in clauses.items()]
    return ParsedDocument(document_name=name, fragments=fragments, clauses={f.clause_id: f for f in fragments}, units=[], functions=[])


def test_new_removed_and_changed_clauses():
    before = make_doc("before", {"1.1": "old", "1.2": "same", "1.3": "removed"})
    after = make_doc("after", {"1.1": "new", "1.2": "same", "1.4": "added"})
    rows = {row.before.clause_id if row.before else row.after.clause_id: row for row in compare_documents(before, after)}
    assert rows["1.1"].status == "Изменена"
    assert rows["1.3"].status == "Удалена / не найдена в ПОСЛЕ"
    assert rows["1.4"].status == "Новая"


def test_fallback_does_not_require_api_key(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    before = make_doc("before", {"1.1": "same"})
    after = make_doc("after", {"1.1": "same"})
    assert compare_documents(before, after)[0].status == "Сохранена"


def test_semantic_fallback_does_not_turn_missing_clause_into_fact_loss(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    before = make_doc("before", {"2.4.1": "проводит проверку системы внутреннего контроля"})
    after = make_doc("after", {"2.4.9": "осуществляет проверку системы внутреннего контроля"})
    rows, candidates = enhance_rows(before, after, compare_documents(before, after), use_ai=False)
    assert candidates == 1
    assert rows[0].status in {"Изменена", "Потенциально перераспределена"}
    assert rows[0].semantic_status != "потенциально потеряна"


def test_ai_source_validation_rejects_unknown_fragment():
    try:
        validate_result({"before_fragment_id": "fake", "after_fragment_ids": ["p1-c2"]}, {"p1-c1"}, {"p1-c2"}, "p1-c1", "p1-c2")
    except ValueError:
        return
    raise AssertionError("invalid AI source was accepted")


def test_semantic_ai_path_uses_mock_and_validates_sources(monkeypatch):
    before = make_doc("before", {"2.4.1": "проверяет систему контроля"})
    after = make_doc("after", {"2.4.9": "осуществляет проверку системы контроля"})
    monkeypatch.setattr(semantic, "call_ai", lambda old, new: {
        "before_fragment_id": old.fragment_id,
        "after_fragment_ids": [new.fragment_id],
        "semantic_status": "смысл сохранен",
        "confidence": 0.93,
        "explanation": "Mock semantic match",
        "requires_human_review": False,
    })
    rows, _ = enhance_rows(before, after, compare_documents(before, after), use_ai=True)
    assert rows[0].analysis_method == "AI"
    assert rows[0].semantic_status == "смысл сохранен"


def test_auth_diagnostic_never_exposes_api_key(monkeypatch):
    fake_key = "sk-fake-secret-123456789"
    monkeypatch.setenv("OPENAI_API_KEY", fake_key)
    diagnostic = semantic._error_info(Exception(f"Incorrect API key provided: {fake_key}"))
    assert fake_key not in diagnostic["message"]
    assert "Authentication error: API key rejected by OpenAI." == diagnostic["message"]


def test_non_auth_diagnostic_redacts_credentials(monkeypatch):
    fake_key = "sk-fake-secret-987654321"
    monkeypatch.setenv("OPENAI_API_KEY", fake_key)
    diagnostic = semantic._error_info(Exception(f"request failed token={fake_key}"))
    assert fake_key not in diagnostic["message"]
    assert "[REDACTED]" in diagnostic["message"]


def test_semantic_status_and_result_type_consistency():
    base = {"before_fragment_id": "p1-c1", "after_fragment_ids": ["p1-c2"], "confidence": 0.95}
    saved = semantic.validate_result({**base, "semantic_status": "потенциально потеряна", "explanation": "no change in meaning"}, {"p1-c1"}, {"p1-c2"}, "p1-c1", "p1-c2")
    assert saved["semantic_status"] == "смысл сохранен"  # contradictory explanation is resolved safely
    assert saved["result_type"] == "FACT"
    assert saved["requires_human_review"] is False
    same = semantic.validate_result({**base, "semantic_status": "смысл сохранен", "explanation": "no change in meaning"}, {"p1-c1"}, {"p1-c2"}, "p1-c1", "p1-c2")
    assert same["result_type"] == "FACT"
    assert same["requires_human_review"] is False
    expanded = semantic.validate_result({**base, "semantic_status": "функция расширена", "explanation": "AFTER expands the scope"}, {"p1-c1"}, {"p1-c2"}, "p1-c1", "p1-c2")
    assert expanded["result_type"] == "FACT"
    uncertain = semantic.validate_result({**base, "semantic_status": "недостаточно данных / требуется проверка человеком", "explanation": "ambiguous"}, {"p1-c1"}, {"p1-c2"}, "p1-c1", "p1-c2")
    assert uncertain["result_type"] == "INFERENCE"
    assert uncertain["requires_human_review"] is True


def test_high_confidence_semantic_consistency_clears_false_review():
    normalized = semantic.validate_result({
        "before_fragment_id": "p1-c1", "after_fragment_ids": ["p1-c2"],
        "semantic_status": "недостаточно данных / требуется проверка человеком",
        "confidence": 0.95,
        "explanation": "The action is semantically consistent and the responsibility is maintained.",
        "requires_human_review": True,
    }, {"p1-c1"}, {"p1-c2"}, "p1-c1", "p1-c2")
    assert normalized["semantic_status"] == "смысл сохранен"
    assert normalized["result_type"] == "FACT"
    assert normalized["requires_human_review"] is False


def test_recommendations_use_evidence_and_skip_saved_meaning():
    before = make_doc("before", {"1.1": "same"})
    after = make_doc("after", {"1.1": "same"})
    row = compare_documents(before, after)[0]
    row.semantic_status = "смысл сохранен"
    row.confidence = 0.95
    row.requires_human_review = False
    assert build_recommendations([row]) == []


def test_review_structural_change_creates_evidenced_recommendation():
    before = make_doc("before", {"3.5": "old responsibility"})
    after = make_doc("after", {"3.5": "new responsibility"})
    row = compare_documents(before, after)[0]
    row.semantic_status = "функция существенно изменена"
    row.requires_human_review = True
    recommendations = build_recommendations([row])
    assert recommendations and recommendations[0]["source"] == row.before.fragment_id


def test_expert_rejection_preserves_ai_result_separately():
    before = make_doc("before", {"3.5": "old responsibility"})
    after = make_doc("after", {"3.5": "new responsibility"})
    row = compare_documents(before, after)[0]
    row.semantic_status = "функция существенно изменена"
    row.requires_human_review = True
    key = review_key(row)
    state = {key: {"ai_result": row.model_dump(), "decision": "Отклонено", "comment": "Не подтверждено"}}
    assert row.semantic_status == state[key]["ai_result"]["semantic_status"]
    assert expert_conclusion([row], state)[0]["status"] == "Отклонено экспертом"


def test_priority_risk_vs_saved_fact():
    before = make_doc("before", {"1.1": "same"})
    after = make_doc("after", {"1.1": "same"})
    fact = compare_documents(before, after)[0]
    fact.semantic_status = "смысл сохранен"
    fact.result_type = "FACT"
    risk = fact.model_copy(update={"semantic_status": "потенциально потеряна", "result_type": "RISK_FLAG", "requires_human_review": True})
    assert review_priority(risk)["priority"] == "Проверить в первую очередь"
    assert review_priority(fact)["priority"] == "Информационно"


def test_ui_source_formatter_uses_original_name_and_safe_diff():
    before = make_doc("tmp123.pdf", {"1.1": "old"})
    before = restore_uploaded_name(before, "Положение.pdf")
    fragment = before.fragments[0]
    assert "tmp123" not in source_label(fragment, {fragment.fragment_id: fragment})
    assert "Положение.pdf" in source_label(fragment, {fragment.fragment_id: fragment})
    removed, added = diff_html("old <tag>", "new <tag>")
    assert "diff-removed" in removed and "diff-added" in added
    assert "&lt;tag&gt;" in removed


def test_decision_workflow_keeps_ai_result_immutable():
    before = make_doc("before", {"2.4.1": "old"})
    after = make_doc("after", {"2.4.1": "new"})
    row = compare_documents(before, after)[0]
    row.semantic_status = "потенциально перераспределена"
    record = decision_record(row, russian_recommendation(row))
    original = record["ai_result"].copy()
    updated = apply_employee_decision(record, "Изменить решение", "Согласовано иначе", "Изменено")
    assert updated["ai_result"] == original
    assert updated["effective_decision"] == "Изменено"
    assert updated["source_fragment_ids"]


def test_rejected_recommendation_does_not_resolve_finding():
    record = {"ai_classification": "потенциально потеряна", "effective_decision": "Отклонить рекомендацию", "employee_decision": "Отклонить рекомендацию"}
    result = revalidate_records([record])
    assert result["resolved"] == 0


def test_accepted_decision_resolves_affected_finding():
    record = {"ai_classification": "потенциально потеряна", "effective_decision": "Принять рекомендацию AI", "employee_decision": "Принять рекомендацию AI"}
    result = revalidate_records([record])
    assert result["resolved"] == 1


def test_deferred_decision_is_residual_and_recommendation_is_russian():
    record = {"ai_classification": "потенциально перераспределена", "effective_decision": "Отложить на дополнительную проверку"}
    result = revalidate_records([record])
    assert result["needs_review"] == 1
    assert "Проверить" in russian_recommendation(type("Row", (), {"semantic_status": "потенциально перераспределена", "status": ""})())
    assert "требует проверки" in russian_impact(type("Row", (), {"semantic_status": "потенциально перераспределена", "requires_human_review": True})())


def test_analysis_state_and_human_decision_survive_reconstruction():
    class State(dict):
        __getattr__ = dict.get
        __setattr__ = dict.__setitem__
    state = State()
    before = make_doc("before.pdf", {"1.1": "same"})
    after = make_doc("after.pdf", {"1.1": "same"})
    save_analysis_state(state, before=before, after=after, rows=[], structure=[], duplicates=[], table=None, recommendations=[], text="result")
    save_human_decision(state, "p1-c1", "Принять рекомендацию AI", "ok", "Сохранить функцию")
    rebuilt = get_analysis_state(state)
    assert state["analysis_completed"] is True
    assert rebuilt["before"].document_name == "before.pdf"
    assert state["human_decisions"]["p1-c1"]["employee_decision"] == "Принять рекомендацию AI"
    save_human_decision(state, "p1-c2", "Отклонить рекомендацию", "no")
    assert state["human_decisions"]["p1-c1"]["employee_decision"] != state["human_decisions"]["p1-c2"]["employee_decision"]


def test_source_label_accepts_function_record_without_attribute_error():
    from app import source_label
    parsed = make_doc("original.pdf", {"1.1": "same"})
    row = compare_documents(parsed, parsed)[0]
    assert "original.pdf" in source_label(row.before, {f.fragment_id: f for f in parsed.fragments})


def test_management_conclusion_is_preliminary_and_russian():
    parsed = make_doc("before.pdf", {"1.1": "same"})
    row = compare_documents(parsed, parsed)[0]
    row.semantic_status = "потенциально потеряна"
    row.requires_human_review = True
    text = management_conclusion([row], [], [], {})
    assert "ПРЕДВАРИТЕЛЬНОЕ" in text
    assert "Окончательное заключение формируется" in text
    assert "рисков нет" not in text.lower()


def test_repeated_recommendations_are_grouped_with_contiguous_clauses():
    before = make_doc("before.pdf", {f"3.{number}": f"old responsibility {number}" for number in range(4, 10)})
    after = make_doc("after.pdf", {f"3.{number}": f"new responsibility {number}" for number in range(4, 10)})
    rows = compare_documents(before, after)
    for row in rows:
        row.requires_human_review = True
    recommendations = build_recommendations(rows)
    grouped = group_recommendations(recommendations, rows)
    assert len(recommendations) == 6
    assert len(grouped) == 1
    assert grouped[0]["affected_clauses"] == "3.4–3.9"
    assert grouped[0]["action"] == "Проверить корректность закрепления ответственности между подразделениями."
    assert grouped[0]["status"] == "Ожидает экспертной проверки"


def test_non_contiguous_clause_ids_are_not_merged_into_false_range():
    before = make_doc("before.pdf", {"3.4": "old 3.4", "3.7": "old 3.7", "5.2": "old 5.2"})
    after = make_doc("after.pdf", {"3.4": "new 3.4", "3.7": "new 3.7", "5.2": "new 5.2"})
    rows = compare_documents(before, after)
    recommendations = [{
        "action": "Проверить корректность закрепления ответственности за подразделением.",
        "source": row.before.fragment_id,
    } for row in rows]
    grouped = group_recommendations(recommendations, rows)
    assert grouped[0]["affected_clauses"] == "3.4, 3.7, 5.2"
    assert "3.4–3.7" not in grouped[0]["affected_clauses"]


def test_management_conclusion_hides_raw_ai_explanation_and_unfinished_review_status():
    before = make_doc("before.pdf", {"3.4": "old responsibility"})
    after = make_doc("after.pdf", {"3.4": "new responsibility"})
    row = compare_documents(before, after)[0]
    row.requires_human_review = True
    row.explanation = "The after fragment changes the responsibility and requires review."
    text = management_conclusion([row], [], build_recommendations([row]), {})
    assert "The after fragment" not in text
    assert "Заключение после экспертной проверки" not in text
    assert "Экспертная проверка ещё не начата." in text


def test_management_conclusion_final_state_uses_structured_human_decision():
    before = make_doc("before.pdf", {"3.4": "old responsibility"})
    after = make_doc("after.pdf", {"3.4": "new responsibility"})
    row = compare_documents(before, after)[0]
    row.requires_human_review = True
    row.explanation = "The raw English explanation must never be rendered."
    decisions = {review_key(row): {"employee_decision": "Принять рекомендацию AI"}}
    text = management_conclusion([row], [], build_recommendations([row]), decisions)
    assert "ЗАКЛЮЧЕНИЕ С УЧЁТОМ ЭКСПЕРТНОЙ ПРОВЕРКИ" in text
    assert "The raw English explanation" not in text
    assert "Принято рекомендаций: 1" in text
    assert "Подтверждённые выводы по пунктам: 3.4." in text


def test_format_clause_ids_only_ranges_contiguous_siblings():
    assert format_clause_ids(["3.4", "3.5", "3.6"]) == "3.4–3.6"
    assert format_clause_ids(["3.4", "3.7", "5.2"]) == "3.4, 3.7, 5.2"


def test_docx_parser_traceability(tmp_path):
    from docx import Document
    document = Document()
    document.add_heading("Положение", level=1)
    document.add_paragraph("2.4.1. Проверяет систему внутреннего контроля")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "3.4."
    table.cell(0, 1).text = "Новое подразделение"
    path = tmp_path / "sample.docx"
    document.save(path)
    parsed = parse_docx(path)
    assert parsed.clauses["2.4.1"].source_locator == "paragraph 2"
    assert any("table 1, row 1" in f.source_locator for f in parsed.fragments)
    assert all(f.fragment_id for f in parsed.fragments)


def test_xlsx_parser_traceability_and_analyzer_compatibility(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Функции"
    sheet.append(["2.4.1", "Проверяет систему контроля"])
    sheet.append(["Подразделение", "Департамент анализа"])
    path = tmp_path / "sample.xlsx"
    workbook.save(path)
    parsed = parse_xlsx(path)
    assert parsed.clauses["2.4.1"].source_locator.startswith("sheet Функции, row 1")
    assert all(f.fragment_id for f in parsed.fragments)
    other = parse_xlsx(path)
    assert compare_documents(parsed, other)[0].status == "Сохранена"


def test_unsupported_xls_is_explicit(tmp_path):
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"not parsed")
    try:
        parse_document(path)
    except ValueError as exc:
        assert ".xls" in str(exc)
    else:
        raise AssertionError(".xls should be rejected explicitly")
